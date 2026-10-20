"""Compile approved chapters into distributable documents (.docx, .txt)."""
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable, Mapping, Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|_[^_]+_)")
BULLET_RE = re.compile(r"^\s*[-*•]\s+(.*)")
NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+(.*)")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)")

PathLike = Union[str, Path]


def safe_filename(title: str, extension: str) -> str:
    stem = re.sub(r"[^\w\- ]+", "", title).strip().replace(" ", "_")[:80] or "book"
    return f"{stem}.{extension.lstrip('.')}"


def _add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** and *italic* markers."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_markdown(doc, markdown: str) -> None:
    """Render a small, safe subset of Markdown into the document."""
    buffer: list[str] = []

    def flush():
        if buffer:
            _add_inline(doc.add_paragraph(), " ".join(s.strip() for s in buffer))
            buffer.clear()

    for raw in (markdown or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        if heading := HEADING_RE.match(line):
            flush()
            # Chapter title is level 1; "#"/"##" sections are level 2, "###" level 3, ...
            level = min(max(len(heading.group(1)), 2), 4)
            doc.add_heading(heading.group(2).strip(" #*"), level=level)
        elif bullet := BULLET_RE.match(line):
            flush()
            _add_inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered := NUMBERED_RE.match(line):
            flush()
            _add_inline(doc.add_paragraph(style="List Number"), numbered.group(1))
        elif line.lstrip().startswith(">"):
            flush()
            p = doc.add_paragraph(style="Quote")
            _add_inline(p, line.lstrip("> ").strip())
        elif line.strip() in {"---", "***", "___"}:
            flush()
        else:
            buffer.append(line)
    flush()


def _add_page_number_footer(section) -> None:
    """Centered 'page N' field in the footer (Word renders the live page number)."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)


class BookCompiler:
    @staticmethod
    def compile_to_docx(title: str, chapters: Iterable[Mapping], output_path: PathLike) -> Path:
        chapters = sorted(chapters, key=lambda c: c["chapter_number"])
        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = "Georgia"
        normal.font.size = Pt(11.5)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.25

        section = doc.sections[0]
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
        _add_page_number_footer(section)

        # Title page
        for _ in range(6):
            doc.add_paragraph()
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_paragraph(f"Compiled {datetime.now():%B %d, %Y}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].italic = True
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        # Table of contents
        doc.add_heading("Contents", level=1)
        for ch in chapters:
            doc.add_paragraph(f"Chapter {ch['chapter_number']}: {ch['title']}", style="List Number")
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        # Chapters
        for index, ch in enumerate(chapters):
            doc.add_heading(f"Chapter {ch['chapter_number']}: {ch['title']}", level=1)
            _add_markdown(doc, ch.get("content") or "")
            if index < len(chapters) - 1:
                doc.add_page_break()

        doc.core_properties.title = title
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

    @staticmethod
    def compile_to_txt(title: str, chapters: Iterable[Mapping], output_path: PathLike) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8") as f:
            f.write(f"{title.upper()}\n\n")
            for ch in sorted(chapters, key=lambda c: c["chapter_number"]):
                f.write(f"Chapter {ch['chapter_number']}: {ch['title']}\n")
                f.write("-" * 40 + "\n")
                f.write((ch.get("content") or "") + "\n\n")
        return output_path


compiler = BookCompiler()
